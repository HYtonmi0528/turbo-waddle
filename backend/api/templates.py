import uuid, os, datetime
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession
from database import get_db, config
from models import UserTemplate
from dependencies import get_current_user

router = APIRouter(prefix="/api", tags=["模板"])

@router.get("/mytemplates")
async def list_templates(user=Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    result = await db.execute(
        select(UserTemplate).where(UserTemplate.owner_id == user.id).order_by(UserTemplate.updated_at.desc())
    )
    templates = [{
        "id": t.id, "name": t.name, "type": t.type, "description": t.description,
        "originalName": t.original_name, "structure": t.structure_json or {},
        "mappings": t.mappings_json or [], "isShared": t.is_shared,
        "createdAt": t.created_at.isoformat() if t.created_at else None
    } for t in result.scalars().all()]
    return {"templates": templates}

@router.get("/templates/shared")
async def list_shared(user=Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    from models import User
    result = await db.execute(
        select(UserTemplate, User.display_name).join(User, User.id == UserTemplate.owner_id)
        .where(UserTemplate.is_shared == True, UserTemplate.owner_id != user.id)
        .order_by(UserTemplate.updated_at.desc())
    )
    items = []
    for t, name in result.all():
        items.append({
            "id": t.id, "name": t.name, "type": t.type, "description": t.description,
            "structure": t.structure_json, "mappings": t.mappings_json, "ownerName": name
        })
    return items

@router.delete("/mytemplates/{template_id}")
async def del_template(template_id: str, user=Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    t = (await db.execute(select(UserTemplate).where(UserTemplate.id == template_id, UserTemplate.owner_id == user.id))).scalar_one_or_none()
    if not t: raise HTTPException(404, "模板不存在")
    if t.storage_path and os.path.exists(t.storage_path):
        try: os.unlink(t.storage_path)
        except: pass
    await db.delete(t)
    await db.commit()
    return {"ok": True}

@router.get("/templates/{template_id}/structure")
async def get_structure(template_id: str, db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    t = (await db.execute(select(UserTemplate).where(UserTemplate.id == template_id))).scalar_one_or_none()
    return t.structure_json if t else {}

@router.get("/templates/{template_id}/mappings")
async def get_mappings(template_id: str, db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    t = (await db.execute(select(UserTemplate).where(UserTemplate.id == template_id))).scalar_one_or_none()
    return t.mappings_json if t else []

@router.put("/templates/{template_id}/mappings")
async def put_mappings(template_id: str, mappings: list, db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    await db.execute(update(UserTemplate).where(UserTemplate.id == template_id).values(mappings_json=mappings, updated_at=datetime.datetime.utcnow()))
    await db.commit()
    return {"ok": True}

@router.put("/templates/{template_id}/structure")
async def put_structure(template_id: str, structure: dict, db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    await db.execute(update(UserTemplate).where(UserTemplate.id == template_id).values(structure_json=structure, updated_at=datetime.datetime.utcnow()))
    await db.commit()
    return {"ok": True}

@router.post("/templates/{template_id}/duplicate")
async def dup_template(template_id: str, db: AsyncSession = Depends(get_db), user=Depends(get_current_user)):
    orig = (await db.execute(select(UserTemplate).where(UserTemplate.id == template_id))).scalar_one_or_none()
    if not orig: raise HTTPException(404, "模板不存在")
    new_id = str(uuid.uuid4())
    t = UserTemplate(id=new_id, owner_id=user.id, name=orig.name+" (副本)", type=orig.type,
                     description=orig.description, structure_json=orig.structure_json,
                     mappings_json=orig.mappings_json, storage_path=orig.storage_path)
    db.add(t); await db.commit()
    return {"id": new_id}

@router.post("/templates/import")
async def import_template(file: UploadFile = File(...), user=Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    import openpyxl
    tid = str(uuid.uuid4())
    tpl_dir = os.path.join(config.get("storageDir", "./server-data/files"), "templates")
    os.makedirs(tpl_dir, exist_ok=True)
    sp = os.path.join(tpl_dir, f"{tid}.xlsx")
    content = await file.read()
    with open(sp, "wb") as f: f.write(content)

    try:
        wb = openpyxl.load_workbook(sp)
        ws = wb.active
        fields = []
        for col in range(1, ws.max_column + 1):
            val = str(ws.cell(1, col).value or "").strip()
            if val: fields.append(val)
        structure = {"headerRow": 1, "columns": [{"index": i+1, "name": f, "field": f, "type": "text", "width": 120} for i, f in enumerate(fields)]}
        mappings = [{"templateField": f, "systemField": f} for f in fields]
    except:
        structure = {}
        mappings = []

    name = os.path.splitext(file.filename)[0] if file.filename else "模板"
    t = UserTemplate(id=tid, owner_id=user.id, name=name, type="通用",
                     original_name=file.filename, structure_json=structure,
                     mappings_json=mappings, storage_path=sp)
    db.add(t); await db.commit()
    return {"id": tid, "name": name, "fields": fields, "structure": structure}
