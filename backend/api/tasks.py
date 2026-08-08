import uuid, datetime, json, os
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlalchemy import select, func, and_
from sqlalchemy.ext.asyncio import AsyncSession
from pydantic import BaseModel
from database import get_db, AsyncSessionLocal
from models import RfqTask, RfqItem, Notification, User, TaskComment
from dependencies import get_current_user, require_role

router = APIRouter(prefix="/api/tasks", tags=["任务"])

@router.get("")
async def list_tasks(db: AsyncSession = Depends(get_db), user: User = Depends(get_current_user)):
    result = await db.execute(
        select(RfqTask).order_by(RfqTask.updated_at.desc())
    )
    tasks = result.scalars().all()
    return {
        "tasks": [{
            "id": t.id, "taskNo": t.task_no, "title": t.title, "requester": t.requester,
            "country": t.country, "clientName": t.client_name, "requestDate": t.request_date.isoformat() if t.request_date else None,
            "deadline": t.deadline.isoformat() if t.deadline else None, "importType": t.import_type,
            "deliveryType": t.delivery_type, "paymentType": t.payment_type, "status": t.status,
            "assignedUserIds": t.assigned_user_ids or [], "currentVersion": t.current_version,
            "createdAt": t.created_at.isoformat() if t.created_at else None, "updatedAt": t.updated_at.isoformat() if t.updated_at else None
        } for t in tasks]
    }

@router.get("/stats")
async def stats(db: AsyncSession = Depends(get_db), user: User = Depends(get_current_user)):
    result = await db.execute(select(RfqTask.status, func.count(RfqTask.id)).group_by(RfqTask.status))
    by_status = [{"status": r[0], "count": r[1]} for r in result.all()]
    total = await db.scalar(select(func.count()).select_from(RfqItem))
    filled = await db.scalar(select(func.count()).select_from(RfqItem).where(RfqItem.fob_usd.isnot(None)))
    return {"byStatus": by_status, "totalItems": total or 0, "filledItems": filled or 0}

@router.get("/{task_id}")
async def get_task(task_id: str, db: AsyncSession = Depends(get_db), user: User = Depends(get_current_user)):
    t = await db.scalar(select(RfqTask).where(RfqTask.id == task_id))
    if not t: raise HTTPException(404, "任务不存在")
    items_result = await db.execute(select(RfqItem).where(RfqItem.task_id == task_id).order_by(RfqItem.line_no))
    items = items_result.scalars().all()
    return {
        "task": {
            "id": t.id, "taskNo": t.task_no, "title": t.title, "requester": t.requester,
            "country": t.country, "clientName": t.client_name, "requestDate": t.request_date.isoformat() if t.request_date else None,
            "deadline": t.deadline.isoformat() if t.deadline else None, "importType": t.import_type,
            "deliveryType": t.delivery_type, "paymentType": t.payment_type, "status": t.status,
            "sourceOriginalName": t.source_original_name, "assignedUserIds": t.assigned_user_ids or [],
            "currentVersion": t.current_version, "sourceStoragePath": t.source_storage_path
        },
        "items": [{
            "id": i.id, "lineNo": i.line_no, "description": i.description, "quantity": float(i.quantity) if i.quantity else None,
            "unit": i.unit, "productCode": i.product_code, "fobUsd": float(i.fob_usd) if i.fob_usd else None,
            "totalRmb": float(i.total_rmb) if i.total_rmb else None, "exchangeRate": float(i.exchange_rate) if i.exchange_rate else None,
            "invoiceType": i.invoice_type, "selectedSupplier": i.selected_supplier,
            "remarks": i.remarks, "attachments": i.attachments_json or [],
            "source": i.source_json or {}, "rowVersion": i.row_version,
            "updatedAt": i.updated_at.isoformat() if i.updated_at else None
        } for i in items]
    }

class ImportMeta(BaseModel):
    title: str = ""
    taskNo: str = ""
    deadline: str = ""
    assignedUserIds: list = []

@router.post("/import")
async def import_task(file: UploadFile = File(...), user: User = Depends(require_role("admin", "manager")), db: AsyncSession = Depends(get_db)):
    import openpyxl, datetime, io
    config_path = os.path.join(os.path.dirname(__file__), '..', '..', 'server-data', 'files')
    os.makedirs(os.path.join(config_path, 'tasks'), exist_ok=True)
    tid = str(uuid.uuid4())
    file_path = os.path.join(config_path, 'tasks', tid)
    os.makedirs(file_path, exist_ok=True)
    content = await file.read()
    safe_name = file.filename.replace(" ", "_") if file.filename else "import.xlsx"
    dest = os.path.join(file_path, safe_name)
    with open(dest, "wb") as f: f.write(content)

    wb = openpyxl.load_workbook(dest, data_only=True)
    # Find best sheet
    best_ws = None; best_hits = 0; best_hr = 1
    FIELD_ALIASES = {
        "description": ["descripción","descripcion","description","描述","产品描述","品名","product description"],
        "quantity": ["cantidad","cant.","quantity","数量","qty","quantities"],
        "unit": ["unidad","unit","单位","u/m"],
        "code": ["código","codigo","code","ref","reference","referencia","型号","代码","product code"],
        "ltc": ["ltc","hs code","arancel","关税"]
    }
    for ws in wb.worksheets:
        hits = 0; header_row = 1
        for r in range(1, min(41, ws.max_row + 1)):
            row_hits = 0
            row_vals = [str(ws.cell(r, c).value or "").lower().strip() for c in range(1, min(51, ws.max_column + 1))]
            for aliases in FIELD_ALIASES.values():
                for v in row_vals:
                    for a in aliases:
                        if a in v: row_hits += 1; break
            if row_hits >= 3: header_row = r; hits = row_hits; break
        if hits > best_hits: best_hits = hits; best_ws = ws; best_hr = header_row

    if not best_ws: best_ws = wb.active
    ws = best_ws

    # Map columns
    col_map = {}
    for c in range(1, ws.max_column + 1):
        val = str(ws.cell(best_hr, c).value or "").lower().strip()
        for field, aliases in FIELD_ALIASES.items():
            for a in aliases:
                if a in val: col_map[field] = c; break
            if field in col_map: break

    items = []; blanks = 0
    for r in range(best_hr + 1, ws.max_row + 1):
        desc = str(ws.cell(r, col_map.get("description") or 1).value or "").strip()
        if desc: blanks = 0
        else: blanks += 1
        if blanks > 8: break
        if not desc: continue
        raw = str(ws.cell(r, 1).value or "").lower()
        if any(kw in raw for kw in ["total","subtotal","suma","合计"]): break
        items.append({
            "lineNo": len(items) + 1, "description": desc,
            "quantity": ws.cell(r, col_map.get("quantity") or 2).value if col_map.get("quantity") else None,
            "unit": str(ws.cell(r, col_map.get("unit") or 2).value or "") if col_map.get("unit") else "",
            "productCode": str(ws.cell(r, col_map.get("code") or 2).value or "") if col_map.get("code") else "",
            "source": {}
        })

    task = RfqTask(id=tid, task_no=f"RFQ-{datetime.date.today()}-{int(__import__('time').time())%10000}",
                   title=safe_name.replace(".xlsx",""), status="published",
                   source_original_name=safe_name, source_storage_path=dest, created_by=user.id)
    db.add(task)
    for item in items:
        db.add(RfqItem(id=str(uuid.uuid4()), task_id=tid, line_no=item["lineNo"], description=item["description"],
                        quantity=item.get("quantity"), unit=item.get("unit"), product_code=item.get("productCode")))
    await db.commit()
    return {"id": tid, "title": task.title, "itemCount": len(items)}

class UpdateItemReq(BaseModel):
    fobUsd: float = None
    totalRmb: float = None
    exchangeRate: float = None
    invoiceType: str = "special"
    selectedSupplier: str = None
    remarks: str = None
    rowVersion: int

@router.patch("/{task_id}/items/{item_id}")
async def update_item(task_id: str, item_id: str, req: UpdateItemReq, user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    item = await db.scalar(select(RfqItem).where(RfqItem.id == item_id, RfqItem.task_id == task_id))
    if not item: raise HTTPException(404, "产品明细不存在")
    if item.row_version != req.rowVersion: raise HTTPException(409, "该行已被其他员工修改，请刷新")
    item.fob_usd = req.fobUsd
    item.total_rmb = req.totalRmb
    item.exchange_rate = req.exchangeRate
    item.invoice_type = req.invoiceType
    item.selected_supplier = req.selectedSupplier
    item.remarks = req.remarks
    item.row_version += 1
    item.updated_by = user.id
    task = await db.scalar(select(RfqTask).where(RfqTask.id == task_id))
    if task and task.status == "published": task.status = "in_progress"
    await db.commit()
    return {"rowVersion": item.row_version, "fobUsd": item.fob_usd, "totalRmb": item.total_rmb, "updatedAt": item.updated_at.isoformat() if item.updated_at else None}

@router.post("/{task_id}/items/{item_id}/revert")
async def revert_item(task_id: str, item_id: str, user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    item = await db.scalar(select(RfqItem).where(RfqItem.id == item_id, RfqItem.task_id == task_id))
    if not item: raise HTTPException(404)
    item.fob_usd = None; item.total_rmb = None; item.selected_supplier = None; item.remarks = None
    item.row_version += 1; item.updated_by = user.id
    await db.commit()
    return {"ok": True}

@router.get("/{task_id}/comments")
async def get_comments(task_id: str, db: AsyncSession = Depends(get_db), user: User = Depends(get_current_user)):
    result = await db.execute(
        select(TaskComment.content, TaskComment.created_at, User.display_name)
        .join(User, User.id == TaskComment.user_id)
        .where(TaskComment.task_id == task_id).order_by(TaskComment.created_at)
    )
    return {"comments": [{"userName": r[2], "content": r[0], "createdAt": r[1].isoformat()} for r in result.all()]}

@router.post("/{task_id}/comments")
async def add_comment(task_id: str, content: str = Query(...), user: User = Depends(get_current_user), db: AsyncSession = Depends(get_db)):
    c = TaskComment(id=str(uuid.uuid4()), task_id=task_id, user_id=user.id, content=content)
    db.add(c); await db.commit()
    return {"id": c.id, "userName": user.display_name, "content": content, "createdAt": c.created_at.isoformat()}
