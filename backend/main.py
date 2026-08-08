from fastapi import FastAPI, HTTPException, Depends, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, FileResponse
from api import auth, tasks, users, templates, data
from dependencies import get_current_user
from models import User, UserTemplate, RfqTask, Notification, DataEntry, DataDraft
from sqlalchemy import select, update
import os

app = FastAPI(title="LATIC询价协作系统", version="2.1.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

app.include_router(auth.router)
app.include_router(tasks.router)
app.include_router(users.router)
app.include_router(templates.router)
app.include_router(data.router)

@app.get("/api/health")
async def health():
    return {"ok": True, "service": "LATIC API v2", "time": __import__("datetime").datetime.utcnow().isoformat()}

@app.get("/api/setup/status")
async def setup_status():
    from database import AsyncSessionLocal
    from sqlalchemy import select, func
    async with AsyncSessionLocal() as db:
        count = await db.scalar(select(func.count()).select_from(User))
        return {"initialized": count > 0, "canInitialize": True}

@app.get("/api/exchange-rate")
async def exchange_rate(force: str = ""):
    import json
    cache_path = os.path.join(os.path.dirname(__file__), '..', 'server-data', 'exrate.json')
    today = __import__("datetime").date.today().isoformat()
    if force != "1":
        try:
            with open(cache_path) as f: cached = json.load(f)
            if cached.get("date") == today: return cached
        except: pass
    try:
        import urllib.request
        resp = urllib.request.urlopen("https://api.frankfurter.dev/latest?from=USD&to=CNY", timeout=8)
        data = json.loads(resp.read())
        result = {"rate": data["rates"]["CNY"], "date": data["date"], "source": "frankfurter.dev"}
        os.makedirs(os.path.dirname(cache_path), exist_ok=True)
        with open(cache_path, "w") as f: json.dump(result, f)
        return result
    except: return {"rate": 7.25, "date": today, "source": "默认"}

@app.get("/api/search")
async def search(q: str = ""):
    if len(q) < 2: return {"tasks": [], "items": []}
    from database import AsyncSessionLocal
    like = f"%{q}%"
    from sqlalchemy import or_
    from models import RfqItem
    async with AsyncSessionLocal() as db:
        tasks = await db.execute(
            select(RfqTask).where(or_(RfqTask.title.like(like), RfqTask.task_no.like(like), RfqTask.requester.like(like))).limit(20)
        )
        items = await db.execute(
            select(RfqItem.id, RfqItem.line_no, RfqItem.description, RfqItem.selected_supplier, RfqItem.task_id, RfqTask.title, RfqTask.task_no)
            .join(RfqTask, RfqTask.id == RfqItem.task_id)
            .where(or_(RfqItem.description.like(like), RfqItem.product_code.like(like), RfqItem.selected_supplier.like(like))).limit(20)
        )
        return {
            "tasks": [{"id": t.id, "taskNo": t.task_no, "title": t.title, "status": t.status} for t in tasks.scalars()],
            "items": [{"id": r[0], "lineNo": r[1], "description": r[2], "supplier": r[3], "taskId": r[4], "taskTitle": r[5], "taskNo": r[6]} for r in items.all()]
        }

@app.get("/api/notifications")
async def get_notifications(user=Depends(get_current_user)):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        result = await db.execute(select(Notification).where(Notification.user_id == user.id).order_by(Notification.created_at.desc()).limit(50))
        return {"notifications": [{"id": n.id, "title": n.title, "message": n.message, "type": n.type, "taskId": n.task_id, "isRead": n.is_read, "createdAt": n.created_at.isoformat() if n.created_at else None} for n in result.scalars().all()]}

@app.patch("/api/notifications/{nid}/read")
async def read_notif(nid: str, user=Depends(get_current_user)):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        await db.execute(update(Notification).where(Notification.id==nid, Notification.user_id==user.id).values(is_read=True))
        await db.commit()
    return {"ok": True}

@app.get("/api/fields")
async def list_fields(user=Depends(get_current_user)):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        from sqlalchemy import text
        try:
            result = await db.execute(text("SELECT field_key, label, category, data_type FROM custom_system_fields ORDER BY label"))
            return {"fields": [{"key": r[0], "label": r[1], "category": r[2], "dataType": r[3]} for r in result.all()]}
        except: return {"fields": []}

@app.post("/api/fields")
async def add_field(body: dict, user=Depends(get_current_user)):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        from sqlalchemy import text
        await db.execute(text("INSERT INTO custom_system_fields (field_key, label, category, data_type, created_by, created_at) VALUES (:k,:l,:c,:d,:u,NOW()) ON DUPLICATE KEY UPDATE label=VALUES(label)"),
                         {"k": body["key"], "l": body["label"], "c": body.get("category","自定义"), "d": body.get("dataType","text"), "u": user.id})
        await db.commit()
    return {"ok": True}

@app.delete("/api/fields/{key}")
async def del_field(key: str):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        from sqlalchemy import text
        await db.execute(text("DELETE FROM custom_system_fields WHERE field_key=:k"), {"k": key})
        await db.commit()
    return {"ok": True}

@app.get("/api/db/tables")
async def db_tables(user=Depends(get_current_user)):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        from sqlalchemy import text
        result = await db.execute(text("SELECT TABLE_NAME, TABLE_ROWS FROM information_schema.TABLES WHERE TABLE_SCHEMA='latic_rfq' AND TABLE_TYPE='BASE TABLE' ORDER BY TABLE_NAME"))
        return [{"name": r[0], "rowCount": r[1] or 0} for r in result.all()]

@app.get("/api/db/table/{name}")
async def db_table(name: str, page: int = 0, pageSize: int = 100):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        from sqlalchemy import text
        cols = await db.execute(text(f"SELECT COLUMN_NAME FROM information_schema.COLUMNS WHERE TABLE_NAME='{name}' ORDER BY ORDINAL_POSITION"))
        columns = [r[0] for r in cols.all()]
        total = await db.execute(text(f"SELECT COUNT(*) FROM `{name}`"))
        total_val = total.scalar()
        rows = await db.execute(text(f"SELECT * FROM `{name}` LIMIT {pageSize} OFFSET {page*pageSize}"))
        row_list = [dict(zip(columns, row)) for row in rows.all()]
        return {"columns": columns, "rows": row_list, "total": total_val}

@app.post("/api/db/query")
async def db_query(body: dict):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        from sqlalchemy import text
        sql = (body.get("query","")).strip()
        if not sql.upper().startswith("SELECT"): raise HTTPException(400, "仅支持SELECT")
        result = await db.execute(text(sql))
        try:
            columns = list(result.keys())
            rows = [dict(zip(columns, row)) for row in result.all()]
            return {"columns": columns, "rows": rows, "total": len(rows)}
        except: return {"columns": [], "rows": [], "total": 0}

@app.post("/api/excel/generate")
async def excel_generate_post(body: dict, user=Depends(get_current_user)):
    import openpyxl, io, uuid, datetime
    from database import AsyncSessionLocal
    tid = body.get("templateId")
    if not tid: raise HTTPException(400, "请选择模板")
    async with AsyncSessionLocal() as db:
        tpl = (await db.execute(select(UserTemplate).where(UserTemplate.id == tid))).scalar_one_or_none()
        if not tpl or not tpl.storage_path or not os.path.exists(tpl.storage_path):
            raise HTTPException(400, "模板文件不存在")
        wb = openpyxl.load_workbook(tpl.storage_path)
        ws = wb.active
        structure = tpl.structure_json or {}
        hr = structure.get("headerRow", 1) + 1
        batches = body.get("batches", [])
        row_num = hr
        for batch in batches:
            for item in (batch.get("items") or []):
                for col_idx, val in enumerate(item.values(), 1):
                    try: ws.cell(row=row_num, column=col_idx, value=val)
                    except: pass
                row_num += 1
        output = io.BytesIO()
        wb.save(output); output.seek(0)
        storage_dir = os.path.join(config.get("storageDir", "./server-data/files"), "generated")
        os.makedirs(storage_dir, exist_ok=True)
        out_path = os.path.join(storage_dir, f"{uuid.uuid4()}.xlsx")
        with open(out_path, "wb") as f: f.write(output.getvalue())
        output.seek(0)
        entry = DataEntry(id=str(uuid.uuid4()), user_id=user.id, type="excel_generated",
                         data={"templateName": tpl.name, "rows": row_num - hr, "filePath": out_path, "generatedAt": datetime.datetime.utcnow().isoformat()})
        db.add(entry); await db.commit()
        return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment;filename=generated_{datetime.date.today()}.xlsx"})

@app.post("/api/excel/open-existing")
async def excel_open_existing(file: UploadFile = File(...)):
    import openpyxl
    content = await file.read()
    wb = openpyxl.load_workbook(__import__('io').BytesIO(content), data_only=True)
    ws = wb.active
    batches = [{"category": "", "items": []}]
    for r in range(2, ws.max_row + 1):
        first = str(ws.cell(r, 1).value or "").strip()
        if not first and not ws.cell(r, 2).value: continue
        item = {}
        for c in range(1, ws.max_column + 1): item[f"col_{c}"] = ws.cell(r, c).value
        batches[0]["items"].append(item)
    return {"success": True, "batches": batches}

@app.post("/api/tasks/{id}/submit-review")
async def submit_review(id: str, user=Depends(get_current_user)):
    from database import AsyncSessionLocal
    import uuid as _uuid, datetime as _dt
    async with AsyncSessionLocal() as db:
        task = (await db.execute(select(RfqTask).where(RfqTask.id == id))).scalar_one_or_none()
        if not task: raise HTTPException(404)
        task.status = "review"; task.updated_at = _dt.datetime.utcnow()
        admins = (await db.execute(select(User).where(User.role.in_(["admin","manager"]), User.status == "active"))).scalars().all()
        for a in admins:
            db.add(Notification(id=str(_uuid.uuid4()), user_id=a.id, type="review", title="询价任务等待审核", message=f"{user.display_name}提交了任务", task_id=id))
        await db.commit()
    return {"ok": True}

@app.post("/api/tasks/{id}/snapshots")
async def snapshots(id: str, user=Depends(get_current_user)):
    from database import AsyncSessionLocal
    async with AsyncSessionLocal() as db:
        task = (await db.execute(select(RfqTask).where(RfqTask.id == id))).scalar_one_or_none()
        if not task: raise HTTPException(404)
        version = task.current_version or 1
        task.current_version = version + 1; task.status = "submitted"; task.updated_at = __import__("datetime").datetime.utcnow()
        await db.commit()
        return {"ok": True, "version": version + 1}

@app.post("/api/tasks/{tid}/items/{iid}/attachments")
async def upload_attachment(tid: str, iid: str, file: UploadFile = File(...)):
    store = os.path.join(os.environ.get("STORAGE_DIR", "./server-data/files"), "tasks", tid, "attachments")
    os.makedirs(store, exist_ok=True)
    fpath = os.path.join(store, f"{__import__('uuid').uuid4()}-{file.filename}")
    with open(fpath, "wb") as f:
        while chunk := await file.read(65536): f.write(chunk)
    return {"ok": True, "name": file.filename, "size": os.path.getsize(fpath)}

@app.get("/api/tasks/{tid}/items/{iid}/attachments/{aid}")
async def get_attachment(tid: str, iid: str, aid: str):
    store = os.path.join(os.environ.get("STORAGE_DIR", "./server-data/files"), "tasks", tid, "attachments")
    if not os.path.exists(store): raise HTTPException(404)
    files = [f for f in os.listdir(store) if f.startswith(aid) or aid in f]
    if not files: raise HTTPException(404)
    return FileResponse(os.path.join(store, files[0]))

@app.on_event("startup")
async def startup():
    from database import engine, Base
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

frontend_dir = os.path.join(os.path.dirname(__file__), "..", "frontend", "dist")
if os.path.exists(frontend_dir):
    from fastapi.staticfiles import StaticFiles
    app.mount("/assets", StaticFiles(directory=os.path.join(frontend_dir, "assets")), name="assets")
    @app.get("/logo.png")
    async def serve_logo(): return FileResponse(os.path.join(frontend_dir, "logo.png"))
    @app.get("/{full_path:path}")
    async def serve_spa(full_path: str = ""):
        index_path = os.path.join(frontend_dir, "index.html")
        if os.path.exists(index_path): return FileResponse(index_path)
    @app.get("/")
    async def serve_root(): return FileResponse(os.path.join(frontend_dir, "index.html"))
